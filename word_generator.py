from docx import Document
import os

def generate_word_document(topic: str, content: str) -> str:
    """
    Génère un document Word structuré avec un titre et le contenu fourni.
    """
    document = Document()
    document.add_heading(f"Rapport : {topic}", 0)
    document.add_paragraph("Document généré automatiquement à partir des recherches Web.\n")

    document.add_heading("Contenu", level=1)
    document.add_paragraph(content)

    folder = "generated_docs"
    os.makedirs(folder, exist_ok=True)
    filename = f"{topic.strip().replace(' ', '_')}.docx"
    path = os.path.join(folder, filename)
    document.save(path)

    return path
